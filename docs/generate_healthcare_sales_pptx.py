"""
Healthcare product briefing — updated to the shipped Foundary baseline (Sep 2026).
Run: python3 docs/generate_healthcare_sales_pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
INK = RGBColor(0x1B, 0x24, 0x33)
MUTED = RGBColor(0x5B, 0x67, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xF4, 0xF6, 0xFA)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GOLD = RGBColor(0xB4, 0x53, 0x09)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Foundary-Healthcare-AI-Security-Brief.pptx"
TOTAL = 18

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp


def text(s, content, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = content
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def bullets(s, items, x, y, w, h, size=15):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = "•  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = INK
        run.font.name = "Calibri"
        p.space_after = Pt(7)
    return box


def footer(s, page):
    text(s, "Foundary  |  Confidential  |  Healthcare AI Security  |  Product baseline", 0.5, 7.15, 10, 0.25, 10, False, MUTED)
    text(s, f"{page}  /  {TOTAL}", 11.6, 7.15, 1.3, 0.25, 10, False, MUTED, PP_ALIGN.RIGHT)


def header_bar(s, kicker, title):
    rect(s, 0, 0, 13.333, 7.5, PALE)
    rect(s, 0, 0, 0.12, 7.5, BLUE)
    text(s, kicker.upper(), 0.5, 0.22, 12, 0.28, 11, True, BLUE)
    text(s, title, 0.5, 0.48, 12.3, 0.55, 24, True, NAVY)


# 1
s = slide()
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 0.14, 7.5, GOLD)
text(s, "CONFIDENTIAL  ·  UPDATED PRODUCT BRIEFING  ·  SEPTEMBER 2026", 0.6, 0.4, 12, 0.3, 12, True, GOLD)
text(s, "Foundary", 0.6, 1.0, 12, 0.55, 18, False, GOLD)
text(s, "Prove the clinical copilot\ncannot be talked into\nhanding over the chart", 0.6, 1.55, 12.2, 2.4, 36, True, WHITE)
text(
    s,
    "Closed-loop AI security on Microsoft Azure AI Foundry — discover, evidence, validate,\n"
    "retest, incident, executive GO / NO-GO. Shipped baseline, not multi-tenant SaaS.",
    0.6,
    4.2,
    12,
    1.1,
    17,
    False,
    RGBColor(0xC5, 0xD0, 0xDE),
)
text(
    s,
    "GitHub: github.com/atulmishra84/Foundry-Baseline-Security-Agents    ·    Project: project-security-agents",
    0.6,
    6.45,
    12,
    0.35,
    13,
    False,
    RGBColor(0x9A, 0xA8, 0xB8),
)
text(s, f"1  /  {TOTAL}", 11.6, 7.15, 1.3, 0.25, 10, False, RGBColor(0x9A, 0xA8, 0xB8), PP_ALIGN.RIGHT)

# 2
header_bar(s := slide(), "The decision", "What a health-system buyer is actually buying")
bullets(
    s,
    [
        "A control plane for AI agents that can search patients, book visits, or draft clinical text — not another SOC seat.",
        "It runs where those agents are built: Azure AI Foundry, with a Container App that executes the loop and writes evidence.",
        "Scope is authorized systems only. We will not red-team Epic/Cerner production. That refusal is a product feature.",
        "You get fail / fix / prove: unprotected demo CONFIRMED T-001; hardened copy BLOCKED; dashboard states NO-GO vs GO.",
        "You do not get 50-hospital SaaS, Entra-only on every outbound call, or a published Power BI workspace until the tenant is licensed.",
    ],
    0.5,
    1.25,
    12.3,
    5.5,
    17,
)
footer(s, 2)

# 3
header_bar(s := slide(), "Why healthcare", "The unit of risk is the agent with a tool, not the website")
pairs = [
    ("HIPAA still applies", "If the model can dump a store, that is an ePHI pathway even when the UI is a chat box."),
    ("Boards will ask for evidence", "“Show the copilot cannot ignore policy.” An annual pentest PDF is not a control proof."),
    ("Azure is already the factory", "IDNs building copilots on Foundry. Security that lives elsewhere becomes shadow IT."),
    ("Allowlist is clinical safety", "A rogue prompt-injection test against the EHR is an incident, not a demo."),
]
for i, (h, b) in enumerate(pairs):
    x, y = 0.5 + (i % 2) * 6.4, 1.25 + (i // 2) * 2.6
    rect(s, x, y, 6.15, 2.4, WHITE)
    text(s, h, x + 0.25, y + 0.25, 5.7, 0.45, 18, True, NAVY)
    text(s, b, x + 0.25, y + 0.85, 5.7, 1.3, 15, False, MUTED)
footer(s, 3)

# 4
header_bar(s := slide(), "The gap", "Known Microsoft and market tools vs this loop")
text(s, "They already own", 0.5, 1.15, 6, 0.3, 15, True, RED)
bullets(
    s,
    [
        "Defender / CSPM — hosts, identities, public storage.",
        "WAF / APIM — HTTP, not instruction override.",
        "SAST — libraries, not null patient_id dumps.",
        "Security Copilot — SOC analyst, not your AI SoT.",
        "Yearly pentest — one report, no retest artifact.",
    ],
    0.5,
    1.5,
    6.0,
    4.9,
    15,
)
text(s, "What Foundary proved on the healthcare demo", 6.9, 1.15, 6, 0.3, 15, True, TEAL)
bullets(
    s,
    [
        "T-001: override returned bulk mock records.",
        "T-002 / T-006: patient_search + admin identity.",
        "SOC incident CRITICAL; containment human-gated.",
        "Retest on hardened copy: T-001 BLOCKED.",
        "Dashboard: NO-GO unprotected / GO hardened set.",
    ],
    6.9,
    1.5,
    5.9,
    4.9,
    15,
)
footer(s, 4)

# 5
header_bar(s := slide(), "The product", "Four Foundry agents, one Source of Truth, one runtime")
cards = [
    ("AI Security Engineer", "Discovers the asset (tools, identities). Scores T-001–T-010. Maps OWASP LLM, MITRE ATLAS, NIST AI RMF. Drafts remediations."),
    ("AI Red Team", "Validates only allowlisted targets. Writes CONFIRMED or BLOCKED. Retests T-001 on the hardened copy. No freelance exploits."),
    ("AI Runtime SOC", "Turns evidence into incidents and lineage. Containment needs an approval file. Graph disable is one demo SP only."),
    ("Orchestrator + runtime", "Playground calls POST /v1/lifecycle via OpenAPI. JSON to blob runtime-output. Light executive HTML for leadership."),
]
for i, (h, b) in enumerate(cards):
    y = 1.15 + i * 1.35
    rect(s, 0.5, y, 12.3, 1.25, WHITE)
    text(s, h, 0.75, y + 0.12, 11.8, 0.32, 17, True, NAVY)
    text(s, b, 0.75, y + 0.5, 11.8, 0.6, 14, False, MUTED)
footer(s, 5)

# 6
header_bar(s := slide(), "Technical architecture", "What is running in the tenant today")
bullets(
    s,
    [
        "Foundry hub hub-ai-security + project project-security-agents (East US). Model: existing gpt-4o — do not create a second deployment.",
        "Container App ca-ai-security-runtime: health, lifecycle?retest=true, artifacts, dashboard (Bearer token).",
        "Loop: discover → assess → remediate → authorized validate → SOC incident → retest T-001.",
        "SoT in blob + git: threats, controls, OWASP 2025, ATLAS, NIST AI RMF. Schemas for every JSON artifact.",
        "Safety: config/production.json allowlist; PHI redacted in evidence; A2A bus stamps which agent wrote which file.",
        "Network: vnet-ai-security with private endpoints for blob, Key Vault, OpenAI. Public access still on until Bastion/VPN.",
        "Code: github.com/atulmishra84/Foundry-Baseline-Security-Agents  ·  CI runs schema, retest, and lifecycle tests.",
    ],
    0.5,
    1.15,
    12.3,
    5.6,
    15,
)
footer(s, 6)

# 7
header_bar(s := slide(), "How to use it", "Operator path — week one")
bullets(
    s,
    [
        "Local: python3 tests/mvp_lifecycle.py then open output/executive-dashboard.html.",
        "Azure: Entra user opens the Foundry project (not the OpenAI resource, not Entra app list).",
        "Playground: select Deployment gpt-4o (already exists). New thread. Type: Call runLifecycle with retest true.",
        "Or: source .env && curl POST $FQDN/v1/lifecycle with Authorization Bearer FOUNDRY_RUNTIME_TOKEN.",
        "Read GO / NO-GO, assets table, findings (Engineer), validation (Red Team), incidents (SOC).",
        "Containment stays off until output/approvals/{incident_id}.json and FOUNDRY_SOC_EXECUTE=true.",
        "Do not paste attack / exploit / hijack into Playground — Azure content filter will block it.",
    ],
    0.5,
    1.15,
    12.3,
    5.6,
    15,
)
footer(s, 7)

# 8
header_bar(s := slide(), "Where the agents live", "Foundry project — not Microsoft Entra")
bullets(
    s,
    [
        "ai-security-engineer, ai-red-team, ai-runtime-soc, ai-security-orchestrator are Assistants in project-security-agents.",
        "They will never appear under Entra → App registrations. That is by design: agents are not users.",
        "In Entra you may see foundary-demo-healthcare-tool (SOC containment SP) and workspace managed identities.",
        "Portal URL: ai.azure.com → Build → Agents → workspace project-security-agents.",
        "Playground needs the existing gpt-4o deployment selected. Creating another Standard gpt-4o hits quota (10/50 used).",
        "Company policy “Entra only” applies to people and workload identities — next increment, not this baseline.",
    ],
    0.5,
    1.15,
    12.3,
    5.6,
    16,
)
footer(s, 8)

# 9
header_bar(s := slide(), "Executive dashboard", "Leadership questions the HTML now answers")
bullets(
    s,
    [
        "Is this a customer/production incident? No — authorized mock only.",
        "Was real PHI leaked? No — redacted ids; no clinical fields on the page.",
        "Are we safe? Unprotected demo failed T-001; hardened copy blocked it.",
        "Ship / no-ship? NO-GO unprotected demo · GO keep the hardened control set.",
        "Did SOC shut something down? No — human approval required.",
        "Who owns next? Engineer remediations, SOC incident, leadership must not promote the open build.",
        "What was not retested? Everything except T-001 (allowlisted validation only).",
        "Entra-only policy met? Not on agent→model and agent→runtime (keys/Bearer still exist).",
        "Which agent wrote this? Handoff table from a2a-messages.json + asset table from discovery.",
    ],
    0.5,
    1.12,
    12.3,
    5.7,
    15,
)
footer(s, 9)

# 10
header_bar(s := slide(), "Proof", "Last closed loop on the authorized healthcare assistant")
metrics = [("8", "Findings"), ("2", "Critical"), ("3", "High"), ("1", "T-001 CONFIRMED"), ("BLOCKED", "Retest"), ("1+", "Incidents")]
for i, (n, l) in enumerate(metrics):
    x = 0.5 + i * 2.1
    rect(s, x, 1.25, 1.95, 1.65, WHITE)
    text(s, n, x, 1.4, 1.95, 0.7, 20, True, NAVY, PP_ALIGN.CENTER)
    text(s, l, x, 2.15, 1.95, 0.5, 12, False, MUTED, PP_ALIGN.CENTER)
text(
    s,
    "Artifacts: asset, finding, remediation, attack, evidence, incident, OWASP/ATLAS/NIST coverage, A2A log, dashboard HTML + Power BI CSV.\n"
    "Runtime also writes the same JSON to blob container runtime-output.",
    0.5,
    3.15,
    12.3,
    1.1,
    15,
    False,
    INK,
)
bullets(
    s,
    [
        "Engineer owns the asset and findings. Red Team owns validation/retest. SOC owns the incident. Orchestrator owns run order.",
        "That provenance is how a CISO defends the report: not “the model said so,” but which agent wrote which file.",
    ],
    0.5,
    4.4,
    12.3,
    2.1,
    15,
)
footer(s, 10)

# 11
header_bar(s := slide(), "Why buy Foundary", "Healthcare-specific reasons this deserves a budget line")
bullets(
    s,
    [
        "Threat language regulators will borrow: OWASP LLM Top 10, ATLAS, NIST AI RMF — one SoT, not a slide.",
        "Built for agents with tools (patient_search), not for a marketing site WAF rule.",
        "Allowlist stops you from turning security into an EHR outage.",
        "Legal/Privacy can file JSON + GO/NO-GO. That is audit evidence.",
        "Microsoft path you already staff: Foundry, Azure RBAC, optional later Security Copilot YAML.",
        "Time-to-answer for “did the copilot leak?” is one lifecycle, not a six-week review.",
    ],
    0.5,
    1.2,
    12.3,
    5.5,
    16,
)
footer(s, 11)

# 12
header_bar(s := slide(), "Why not “we already bought X”", "Keep those products — they do not replace this loop")
rows = [
    ("Security Copilot", "SOC copilot in Microsoft.", "No allowlisted T-001 retest, no AI SoT, no Foundary GO/NO-GO."),
    ("Defender / Purview / DLP", "Tenant, data, endpoint.", "Blind to instruction-override → tool dump. Complementary."),
    ("Lakera / HiddenLayer / prompt firewalls", "Point control at the prompt.", "No remediations, SOC incident, or board decision strip."),
    ("AppSec + annual pentest", "Needed for everything else.", "Wrong unit of work: code paths vs agent + tools."),
    ("Copilot Studio only", "Fast clinician UX.", "No security system of record."),
]
text(s, "Alternative", 0.5, 1.12, 3.1, 0.28, 12, True, MUTED)
text(s, "Keep it for", 3.7, 1.12, 3.6, 0.28, 12, True, MUTED)
text(s, "It does not replace Foundary", 7.5, 1.12, 5.3, 0.28, 12, True, MUTED)
for i, (a, b, c) in enumerate(rows):
    y = 1.45 + i * 1.0
    rect(s, 0.5, y, 12.3, 0.92, WHITE)
    text(s, a, 0.65, y + 0.18, 2.95, 0.6, 13, True, NAVY)
    text(s, b, 3.7, y + 0.18, 3.55, 0.6, 13, False, MUTED)
    text(s, c, 7.5, y + 0.18, 5.1, 0.6, 13, False, INK)
footer(s, 12)

# 13
header_bar(s := slide(), "Recommendation", "Buy this if  /  do not buy this if")
rect(s, 0.5, 1.2, 6.05, 5.35, WHITE)
text(s, "Buy when", 0.75, 1.35, 5.6, 0.4, 18, True, TEAL)
bullets(
    s,
    [
        "Azure-hosted clinical or RCM agents in the next 12 months.",
        "Compliance asked for AI-specific evidence.",
        "You want fail / fix / prove on an authorized replica first.",
        "You will staff a human approver for containment.",
    ],
    0.7,
    1.9,
    5.6,
    4.3,
    14,
)
rect(s, 6.75, 1.2, 6.05, 5.35, WHITE)
text(s, "Do not buy (this baseline) if", 7.0, 1.35, 5.6, 0.4, 18, True, RED)
bullets(
    s,
    [
        "You need multi-tenant SaaS and a rate card this quarter.",
        "Entra-only on every API call is a hard gate for go-live.",
        "You expect us to attack production EHR. We refuse.",
        "A WAF rule and a training slide would have been enough.",
    ],
    6.95,
    1.9,
    5.6,
    4.3,
    14,
)
footer(s, 13)

# 14
header_bar(s := slide(), "Shipped vs parked", "Baseline is frozen; these are not in the current offer")
text(s, "In the baseline (done)", 0.5, 1.15, 6, 0.3, 16, True, TEAL)
bullets(
    s,
    [
        "Closed loop + allowlist + T-001 retest.",
        "Foundry agents + OpenAPI runtime.",
        "Exec dashboard (Q&A, assets, provenance).",
        "Hybrid private endpoints (public still on).",
        "GitHub + CI schema/retest/lifecycle.",
        "Healthcare sales narrative (this deck).",
    ],
    0.5,
    1.5,
    6.0,
    5.0,
    14,
)
text(s, "Parked (not in this sale unless scoped)", 6.9, 1.15, 6, 0.3, 16, True, GOLD)
bullets(
    s,
    [
        "Entra-only OpenAI + runtime JWT.",
        "Disable public network / CA on VNet.",
        "Power BI workspace (tenant 404 today).",
        "Security Copilot publish in their tenant.",
        "Containment from Container App MI.",
        "Local Docker amd64 on Apple Silicon.",
    ],
    6.9,
    1.5,
    5.9,
    5.0,
    14,
)
footer(s, 14)

# 15
header_bar(s := slide(), "Diligence answers", "What a strong CISO will ask")
bullets(
    s,
    [
        "Production SaaS? No. Controlled platform + proven loop. Productionization is a joint SOW.",
        "Attack our EHR? No. Allowlist only. They put a replica on the list and sign rules of engagement.",
        "Entra-only? Portal yes. Outbound keys still exist. Call it a backlog item in the contract if required.",
        "Hallucinated findings? Deterministic tools write JSON; LLM only annotates. Exec numbers come from tools.",
        "PHI in logs? Redaction on. Still agree LAW / App Insights retention with Privacy.",
        "Playground quota? Use existing gpt-4o. Runtime API is the system of record if the UI fights you.",
    ],
    0.5,
    1.15,
    12.3,
    5.6,
    15,
)
footer(s, 15)

# 16
header_bar(s := slide(), "Commercial motion", "How to land it")
bullets(
    s,
    [
        "Offer: Foundry Security Loop + healthcare scenario pack (patient search / scheduling patterns) + operating model.",
        "Price against: delayed digital-front-door go-live + one reportable investigation. Not against $20/user SOC.",
        "Land: one non-prod assistant on the allowlist. Expand: every Foundry agent that gains a tool.",
        "Microsoft attach: Azure consumption. Optional: they upload our Security Copilot YAML themselves.",
        "Success metric week 4: one GO/NO-GO in the CISO packet with A2A provenance attached.",
    ],
    0.5,
    1.2,
    12.3,
    5.5,
    16,
)
footer(s, 16)

# 17
header_bar(s := slide(), "If they say yes", "90 days")
cols = [
    ("Days 1–30", "SOW. Allowlist one non-prod assistant. Entra RBAC on the project. Weekly lifecycle. Dashboard in CISO review."),
    ("Days 31–60", "Optional Entra outbound. Private-link cutover plan. Approval playbook with Compliance. Fabric if they want Power BI."),
    ("Days 61–90", "Second agent. Copilot YAML in their tenant if they operate a SOC there. Decision: productize vs platform service."),
]
for i, (h, b) in enumerate(cols):
    x = 0.5 + i * 4.2
    rect(s, x, 1.25, 4.0, 5.2, WHITE)
    text(s, h, x + 0.25, 1.45, 3.5, 0.45, 20, True, BLUE)
    text(s, b, x + 0.25, 2.15, 3.5, 3.9, 15, False, INK)
footer(s, 17)

# 18
s = slide()
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 0.14, 7.5, GOLD)
text(s, "THE LINE TO LEAVE IN THE ROOM", 0.6, 1.15, 12, 0.35, 13, True, GOLD)
text(
    s,
    "You already bought Microsoft to build\nthe clinical copilot. Foundary is how you\nprove it cannot be instructed to hand over\nthe chart — and that the fix held.",
    0.6,
    1.7,
    12.2,
    3.0,
    26,
    True,
    WHITE,
)
text(
    s,
    "Pick one authorized assistant. Run the loop. Read GO / NO-GO.\n"
    "github.com/atulmishra84/Foundry-Baseline-Security-Agents",
    0.6,
    5.3,
    12,
    0.9,
    16,
    False,
    RGBColor(0xC5, 0xD0, 0xDE),
)
text(s, f"{TOTAL}  /  {TOTAL}", 11.5, 7.15, 1.4, 0.25, 10, False, RGBColor(0x9A, 0xA8, 0xB8), PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
print("slides", len(prs.slides))
